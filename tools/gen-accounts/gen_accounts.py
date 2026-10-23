#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Công cụ tạo hàng loạt tài khoản học sinh TRỰC TIẾP TRÊN HỆ THỐNG ĐANG CHẠY — tài khoản tạo ra
ĐĂNG NHẬP ĐƯỢC NGAY trên website (không chỉ là một danh sách mã số suông).

Mẫu mã đăng nhập:  {khối 2 số}{số thứ tự 3 số}{năm 4 số}
  Khối 10, số thứ tự 001, năm 2026  ->  100012026
  Khối 11, số thứ tự 001, năm 2026  ->  110012026
  Khối 12, số thứ tự 001, năm 2026  ->  120012026

QUY TRÌNH SAU KHI TẠO (đã được dựng sẵn trong hệ thống, không cần làm gì thêm):
  1) Học sinh đăng nhập lần đầu bằng mã trên + mật khẩu mặc định (123456 theo mặc định).
  2) Hệ thống BẮT BUỘC đổi mật khẩu ngay — không thể bỏ qua.
  3) Sau khi đổi mật khẩu, hệ thống hiện tiếp một form BẮT BUỘC điền họ tên + chọn đúng lớp
     (và có thể điền thêm giới tính, ngày sinh, số điện thoại).
  4) Chỉ sau khi điền xong form đó, học sinh mới vào được giao diện chính của ứng dụng.

CÁCH DÙNG NHANH (máy chủ đang chạy, vd. qua CHAY-WEB.bat)
  python3 gen_accounts.py --grades 10,11,12 --count 40 --year 2026

  → Script sẽ hỏi tài khoản + mật khẩu Ban Giám hiệu/quản trị (để được phép tạo tài khoản), mở
    ảnh mã xác nhận (CAPTCHA) ra trình duyệt để bạn đọc và nhập lại, rồi tạo 40 tài khoản THẬT
    cho MỖI khối (10, 11, 12) ngay trên hệ thống. Danh sách mã đăng nhập + mật khẩu xuất ra 3 file
    trong ./output/: accounts.csv, accounts.xlsx, accounts.txt — in ra phát cho học sinh.

TÙY CHỌN THƯỜNG DÙNG
  --grades 10,11,12      Danh sách khối, cách nhau bằng dấu phẩy (mặc định: 10,11,12)
  --count 40             Số lượng tài khoản cần tạo cho MỖI khối (bắt buộc)
  --year 2026            Năm (mặc định: năm hiện tại)
  --start-seq 1          Số thứ tự bắt đầu (mặc định: 1, tức 001) — dùng khi tạo thêm đợt sau
  --password 123456      Mật khẩu ban đầu dùng chung cho đợt này (học sinh bắt buộc đổi khi đăng nhập)
  --url http://localhost:3000   Địa chỉ máy chủ (mặc định: http://localhost:3000)
  --out-dir ./output      Thư mục xuất file danh sách (mặc định: ./output)
  --username / --password-admin   Truyền sẵn tài khoản quản trị thay vì để script hỏi (xem CẢNH BÁO dưới)
  --offline               CHỈ sinh danh sách mã số ra file (không tạo tài khoản thật, không cần máy chủ)

VÍ DỤ KHÁC
  # Tạo thêm 10 tài khoản khối 10, bắt đầu từ số thứ tự 041 (vd. đã tạo 40 tài khoản đầu)
  python3 gen_accounts.py --grades 10 --count 10 --year 2026 --start-seq 41

  # Máy chủ đang chạy ở địa chỉ khác / cổng khác
  python3 gen_accounts.py --grades 10 --count 20 --year 2026 --url http://localhost:8080

CẢNH BÁO
  --username/--password-admin gõ thẳng trên dòng lệnh có thể bị lưu lại trong lịch sử lệnh của máy
  (vd. Terminal/PowerShell history) — chỉ dùng khi thật sự cần (vd. chạy tự động theo lịch), bình
  thường cứ để trống và trả lời khi script hỏi, mật khẩu sẽ KHÔNG hiện ra trên màn hình.

GHI CHÚ
  - Mật khẩu "123456" khớp với mật khẩu mặc định của hệ thống khi Ban Giám hiệu tạo tài khoản mới.
  - Không cần cài thêm thư viện ngoài để TẠO tài khoản (chỉ dùng thư viện có sẵn của Python). Muốn
    có thêm file accounts.xlsx đẹp hơn thì cài `pip install openpyxl` — không có vẫn ra đủ .csv/.txt.
  - Script gọi đúng các API mà trang quản trị của website đang dùng, với đúng tài khoản Ban Giám
    hiệu/quản trị của bạn — không có "cửa sau" nào khác, không đụng trực tiếp vào file dữ liệu.
"""
import argparse
import csv
import datetime
import getpass
import json
import os
import ssl
import sys
import tempfile
import urllib.error
import urllib.request
import webbrowser


def http_json(url, method="GET", body=None, token=None, insecure=False):
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    req.add_header("Accept", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    ctx = None
    if url.startswith("https://") and insecure:
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
    try:
        with urllib.request.urlopen(req, context=ctx, timeout=20) as resp:
            raw = resp.read().decode("utf-8")
            return resp.status, (json.loads(raw) if raw else {})
    except urllib.error.HTTPError as e:
        raw = e.read().decode("utf-8", "replace")
        try:
            return e.code, json.loads(raw)
        except json.JSONDecodeError:
            return e.code, {"error": raw or e.reason}
    except urllib.error.URLError as e:
        print(f"\nLỗi: không kết nối được tới {url}\n  → {e.reason}", file=sys.stderr)
        print("Kiểm tra lại máy chủ đã chạy chưa (CHAY-WEB.bat) và --url đã đúng cổng chưa.", file=sys.stderr)
        sys.exit(1)


def show_captcha_and_ask(svg: str) -> str:
    """Mở ảnh CAPTCHA (SVG) ra trình duyệt mặc định để người dùng đọc, rồi hỏi lại qua terminal."""
    fd, path = tempfile.mkstemp(suffix=".html")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write(
            "<!doctype html><html><head><meta charset='utf-8'>"
            "<title>Mã xác nhận</title></head><body style='font-family:sans-serif;"
            "display:flex;flex-direction:column;align-items:center;justify-content:center;"
            "height:100vh;margin:0;background:#0b1220;color:#fff'>"
            "<p>Đọc mã số trong hình rồi quay lại cửa sổ dòng lệnh (terminal) để nhập:</p>"
            f"<div style='transform:scale(2.2);margin:40px'>{svg}</div>"
            "</body></html>"
        )
    try:
        webbrowser.open("file://" + path)
    except Exception:
        print(f"(Không tự mở được trình duyệt — hãy tự mở file này để xem mã: {path})")
    return input("Nhập mã xác nhận (CAPTCHA) hiện trong ảnh vừa mở: ").strip()


def login_staff(base_url: str, username: str, password: str, insecure: bool) -> str:
    for attempt in range(1, 4):
        status, cap = http_json(base_url + "/api/auth/captcha", insecure=insecure)
        if status != 200:
            print("Không lấy được mã xác nhận từ máy chủ:", cap, file=sys.stderr)
            sys.exit(1)
        code = show_captcha_and_ask(cap["svg"])
        status, res = http_json(
            base_url + "/api/auth/login", "POST",
            {"username": username, "password": password, "captchaId": cap["id"], "captchaCode": code, "portal": "staff", "remember": False},
            insecure=insecure,
        )
        if status == 200:
            return res["token"]
        print(f"Đăng nhập thất bại: {res.get('error', status)}")
        if attempt < 3:
            print("Thử lại...")
    print("\nKhông đăng nhập được sau 3 lần thử. Kiểm tra lại tài khoản/mật khẩu quản trị và thử chạy lại script.", file=sys.stderr)
    sys.exit(1)


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["STT", "Khối", "Tài khoản đăng nhập", "Mật khẩu ban đầu"])
        for r in rows:
            w.writerow([r["stt"], r["khoi"], r["username"], r["password"]])


def write_txt(rows, path):
    with open(path, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(f"{r['username']}\t{r['password']}\n")


def write_xlsx(rows, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("(Bỏ qua .xlsx: chưa cài openpyxl — chạy `pip install openpyxl` nếu cần file này.)")
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Tài khoản"
    ws.append(["STT", "Khối", "Tài khoản đăng nhập", "Mật khẩu ban đầu"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1747C9")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r["stt"], r["khoi"], r["username"], r["password"]])
    for col, w in zip("ABCD", [6, 12, 22, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return True


def export_files(rows, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    csv_path = os.path.join(out_dir, "accounts.csv")
    txt_path = os.path.join(out_dir, "accounts.txt")
    xlsx_path = os.path.join(out_dir, "accounts.xlsx")
    write_csv(rows, csv_path)
    write_txt(rows, txt_path)
    made_xlsx = write_xlsx(rows, xlsx_path)
    print(f"\n  - {csv_path}")
    print(f"  - {txt_path}")
    if made_xlsx:
        print(f"  - {xlsx_path}")
    print("\nMẫu 3 dòng đầu:")
    for r in rows[:3]:
        print(f"  {r['username']}  /  {r['password']}   ({r['khoi']})")
    if len(rows) > 3:
        print("  ...")


def run_online(args):
    base_url = args.url.rstrip("/")
    username = args.username or input("Tài khoản quản trị / Ban Giám hiệu (đăng nhập Cổng Giáo Viên): ").strip()
    password = args.password_admin or getpass.getpass("Mật khẩu: ")
    token = login_staff(base_url, username, password, args.insecure)
    print("Đăng nhập thành công. Đang tạo tài khoản trên hệ thống...\n")

    rows = []
    for grade in args.grade_list:
        status, res = http_json(
            base_url + "/api/admin/users/bulk-generate", "POST",
            {"grade": int(grade), "count": args.count, "year": args.year, "startSeq": args.start_seq, "password": args.password},
            token=token, insecure=args.insecure,
        )
        if status != 200:
            print(f"Lỗi khi tạo tài khoản khối {grade}: {res.get('error', status)}", file=sys.stderr)
            continue
        print(f"Khối {grade}: {res['message']}")
        for item in res["created"]:
            rows.append({"stt": len(rows) + 1, "khoi": f"Khối {grade}", "username": item["id"], "password": item["password"]})

    if not rows:
        print("\nKhông có tài khoản nào được tạo.", file=sys.stderr)
        sys.exit(1)

    print(f"\nĐã tạo THẬT {len(rows)} tài khoản trên hệ thống — các mã này đăng nhập được ngay.")
    print("Học sinh đăng nhập lần đầu sẽ được yêu cầu đổi mật khẩu rồi điền họ tên + chọn lớp trước khi vào ứng dụng.")
    export_files(rows, args.out_dir)


def run_offline(args):
    rows = []
    for grade in args.grade_list:
        for i in range(args.count):
            seq = args.start_seq + i
            username = f"{int(grade):02d}{seq:03d}{args.year}"
            rows.append({"stt": len(rows) + 1, "khoi": f"Khối {grade}", "username": username, "password": args.password})
    print(f"(Chế độ --offline: chỉ sinh danh sách, CHƯA tạo tài khoản thật trên hệ thống)")
    print(f"Đã sinh {len(rows)} mã (chưa đăng nhập được cho tới khi bạn chạy lại KHÔNG có --offline).")
    export_files(rows, args.out_dir)


def main():
    p = argparse.ArgumentParser(description="Tạo hàng loạt tài khoản học sinh — mặc định tạo THẬT trên hệ thống đang chạy.",
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--grades", default="10,11,12", help="Danh sách khối, cách nhau bởi dấu phẩy (vd. 10,11,12)")
    p.add_argument("--count", type=int, required=True, help="Số lượng tài khoản cần tạo cho MỖI khối")
    p.add_argument("--year", type=int, default=datetime.date.today().year, help="Năm (mặc định: năm hiện tại)")
    p.add_argument("--start-seq", type=int, default=1, help="Số thứ tự bắt đầu (mặc định: 1)")
    p.add_argument("--password", default="123456", help="Mật khẩu ban đầu dùng chung cho đợt này (mặc định: 123456)")
    p.add_argument("--url", default="http://localhost:3000", help="Địa chỉ máy chủ (mặc định: http://localhost:3000)")
    p.add_argument("--out-dir", default="./output", help="Thư mục xuất file danh sách (mặc định: ./output)")
    p.add_argument("--username", default=None, help="Tài khoản quản trị (nếu bỏ trống, script sẽ hỏi)")
    p.add_argument("--password-admin", default=None, help="Mật khẩu quản trị (nếu bỏ trống, script sẽ hỏi — khuyến nghị để trống)")
    p.add_argument("--insecure", action="store_true", help="Bỏ qua kiểm tra chứng chỉ HTTPS (chỉ dùng khi máy chủ chạy HTTPS với chứng chỉ tự ký)")
    p.add_argument("--offline", action="store_true", help="Chỉ sinh danh sách mã số ra file, KHÔNG tạo tài khoản thật / không cần máy chủ")
    args = p.parse_args()

    args.grade_list = [g.strip() for g in args.grades.split(",") if g.strip()]
    if not args.grade_list:
        print("Lỗi: cần ít nhất một khối trong --grades", file=sys.stderr)
        sys.exit(1)
    if args.count <= 0:
        print("Lỗi: --count phải lớn hơn 0", file=sys.stderr)
        sys.exit(1)

    if args.offline:
        run_offline(args)
    else:
        run_online(args)


if __name__ == "__main__":
    main()
